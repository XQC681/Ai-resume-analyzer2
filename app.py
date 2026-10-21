"""AI Resume ATS Checker - Streamlit + Gemini Flash.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description,
and get an ATS score with concrete improvement suggestions.
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------- Configuration -----------------------------

DEFAULT_MODEL = "gemini-2.5-flash"
MAX_RESUME_CHARS = 15000
MAX_JD_CHARS = 6000
MAX_FILE_MB = 5

# Weights used to compute the overall score (must add up to 1.0).
CATEGORY_WEIGHTS = {
    "keywords_relevance": 0.30,
    "content_impact": 0.25,
    "formatting_parsability": 0.20,
    "structure_sections": 0.15,
    "readability_clarity": 0.10,
}

CATEGORY_LABELS = {
    "keywords_relevance": "Keywords & Relevance",
    "content_impact": "Content & Impact",
    "formatting_parsability": "Formatting & Parsability",
    "structure_sections": "Structure & Sections",
    "readability_clarity": "Readability & Clarity",
}

PROMPT_TEMPLATE = """You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.

Evaluate the resume below{jd_clause}. Be honest, specific and strict: do not inflate scores.
Base every comment on what is actually in the resume text. Never invent facts about the candidate.

Score each category from 0 to 100:
- keywords_relevance: skills/keywords and how well they match {match_target}
- content_impact: quantified achievements, action verbs, results over duties
- formatting_parsability: whether the text suggests an ATS-friendly layout (standard headings, no tables/columns/graphics artifacts, consistent dates)
- structure_sections: presence and order of Contact, Summary, Experience, Education, Skills, etc.
- readability_clarity: concise bullets, grammar, consistent tense, appropriate length

Return ONLY a JSON object (no markdown, no commentary) with exactly this shape:
{{
  "category_scores": {{
    "keywords_relevance": <int>,
    "content_impact": <int>,
    "formatting_parsability": <int>,
    "structure_sections": <int>,
    "readability_clarity": <int>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<short point>", "..."],
  "missing_keywords": ["<keyword or skill>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "<resume section>", "issue": "<what is wrong>", "suggestion": "<specific fix>"}}
  ],
  "rewrite_examples": [
    {{"original": "<weak bullet copied from the resume>", "improved": "<stronger version, no invented numbers>"}}
  ]
}}

Rules:
- 3 to 6 strengths, 5 to 10 improvements ordered High to Low, 0 to 3 rewrite examples.
- missing_keywords: {kw_rule}
- In rewrite_examples, use placeholders like [X%] or [number] instead of inventing metrics.
{jd_block}
RESUME:
\"\"\"
{resume}
\"\"\"
"""


# ----------------------------- File parsing -----------------------------

def extract_text(file_name: str, data: bytes) -> str:
    """Extract plain text from PDF, DOCX or TXT bytes."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected. Please upload an unlocked copy.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ----------------------------- Prompt + parsing -----------------------------

def build_prompt(resume_text: str, job_description: str = "") -> str:
    resume_text = resume_text[:MAX_RESUME_CHARS]
    jd = (job_description or "").strip()[:MAX_JD_CHARS]
    if jd:
        return PROMPT_TEMPLATE.format(
            jd_clause=" against the job description provided",
            match_target="the job description",
            kw_rule="important keywords/skills from the job description that are absent from the resume (max 15)",
            jd_block=f'\nJOB DESCRIPTION:\n"""\n{jd}\n"""\n',
            resume=resume_text,
        )
    return PROMPT_TEMPLATE.format(
        jd_clause=" for general ATS-readiness (no specific job description was provided)",
        match_target="common expectations for the candidate's apparent field",
        kw_rule="commonly expected keywords/skills for the candidate's apparent field that are absent (max 15)",
        jd_block="",
        resume=resume_text,
    )


def parse_json_response(text: str) -> dict:
    """Parse JSON from a model response, tolerating markdown fences and extra text."""
    if not text or not text.strip():
        raise ValueError("The model returned an empty response.")
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start == -1 or end <= start:
            raise ValueError("Could not find JSON in the model response.")
        data = json.loads(cleaned[start : end + 1])
    if not isinstance(data, dict):
        raise ValueError("Model response was not a JSON object.")
    return data


def _clamp_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalize_result(data: dict) -> dict:
    """Validate the model output and compute the weighted overall score."""
    raw_scores = data.get("category_scores")
    raw_scores = raw_scores if isinstance(raw_scores, dict) else {}
    scores = {key: _clamp_score(raw_scores.get(key, 0)) for key in CATEGORY_WEIGHTS}
    overall = _clamp_score(sum(scores[k] * w for k, w in CATEGORY_WEIGHTS.items()))

    improvements = []
    for item in data.get("improvements", []) if isinstance(data.get("improvements"), list) else []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "Medium")).strip().capitalize()
        if priority not in ("High", "Medium", "Low"):
            priority = "Medium"
        improvements.append(
            {
                "priority": priority,
                "section": str(item.get("section", "General")).strip() or "General",
                "issue": str(item.get("issue", "")).strip(),
                "suggestion": str(item.get("suggestion", "")).strip(),
            }
        )
    order = {"High": 0, "Medium": 1, "Low": 2}
    improvements.sort(key=lambda i: order[i["priority"]])

    rewrites = []
    for item in data.get("rewrite_examples", []) if isinstance(data.get("rewrite_examples"), list) else []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append({"original": str(item["original"]).strip(), "improved": str(item["improved"]).strip()})

    return {
        "overall_score": overall,
        "category_scores": scores,
        "summary": str(data.get("summary", "")).strip(),
        "strengths": _str_list(data.get("strengths")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "improvements": improvements,
        "rewrite_examples": rewrites,
    }


def score_band(score: int) -> str:
    if score >= 80:
        return "Excellent - likely to pass most ATS filters"
    if score >= 65:
        return "Good - a few improvements will help"
    if score >= 50:
        return "Fair - needs noticeable work"
    return "Weak - needs significant improvement"


# ----------------------------- Gemini call -----------------------------

def analyze_resume(client, model: str, resume_text: str, job_description: str = "") -> dict:
    prompt = build_prompt(resume_text, job_description)
    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
    )
    response = client.models.generate_content(model=model, contents=prompt, config=config)
    return normalize_result(parse_json_response(response.text))


def friendly_error(exc: Exception) -> str:
    msg = str(exc)
    low = msg.lower()
    if "api key" in low or "api_key" in low or "permission" in low or "401" in low or "403" in low:
        return "Your Gemini API key looks invalid or lacks permission. Please check it."
    if "429" in low or "quota" in low or "resource_exhausted" in low or "rate" in low:
        return "Gemini rate limit or quota reached. Wait a minute and try again."
    if "404" in low or "not found" in low:
        return "Model not found. Change the model name in the sidebar (e.g. gemini-2.5-flash)."
    if "json" in low or "model returned" in low or "model response" in low:
        return "The AI response could not be read. Please click Analyze again."
    return f"Something went wrong: {msg}"


def get_api_key() -> str:
    """Look for the key in Streamlit secrets, then environment variables."""
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        key = ""
    return key or os.environ.get("GEMINI_API_KEY", "") or os.environ.get("GOOGLE_API_KEY", "")


# ----------------------------- UI -----------------------------

def render_results(result: dict) -> None:
    overall = result["overall_score"]
    st.subheader("Your ATS Score")
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("Overall ATS Score", f"{overall} / 100")
    with col2:
        st.progress(overall / 100)
        st.caption(score_band(overall))

    if result["summary"]:
        st.info(result["summary"])

    st.markdown("#### Score breakdown")
    for key, label in CATEGORY_LABELS.items():
        val = result["category_scores"][key]
        c1, c2 = st.columns([1, 3])
        c1.write(f"**{label}**")
        c2.progress(val / 100, text=f"{val}/100")

    left, right = st.columns(2)
    with left:
        st.markdown("#### Strengths")
        for s in result["strengths"] or ["No strengths reported."]:
            st.markdown(f"- {s}")
    with right:
        st.markdown("#### Missing keywords")
        if result["missing_keywords"]:
            st.markdown(" ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.markdown("None detected.")

    st.markdown("#### Suggested improvements")
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    if not result["improvements"]:
        st.write("No improvements reported.")
    for item in result["improvements"]:
        with st.expander(f"{icons[item['priority']]} {item['priority']} - {item['section']}: {item['issue'][:80]}"):
            st.markdown(f"**Issue:** {item['issue']}")
            st.markdown(f"**Fix:** {item['suggestion']}")

    if result["rewrite_examples"]:
        st.markdown("#### Example rewrites")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex['original']}")
            st.markdown(f"**After:** {ex['improved']}")
            st.divider()

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.write("Upload your resume to get an ATS score and specific ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if api_key:
            st.success("API key loaded from secrets/environment.")
        else:
            api_key = st.text_input("Gemini API key", type="password", help="Get one free at aistudio.google.com")
        model = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.caption("Resumes are sent to the Gemini API for analysis. Don't upload anything you aren't comfortable sharing.")

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)",
        height=150,
        placeholder="Paste the job posting here for a tailored score...",
    )

    if st.button("Analyze resume", type="primary"):
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            return
        if uploaded is None:
            st.error("Please upload a resume first.")
            return
        data = uploaded.getvalue()
        if len(data) > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is larger than {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded.name, data)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(text) < 100:
            st.error("Very little text was found. If this is a scanned/image PDF, an ATS can't read it either - export a text-based PDF or DOCX.")
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                client = genai.Client(api_key=api_key)
                result = analyze_resume(client, model.strip() or DEFAULT_MODEL, text, job_description)
            st.session_state["result"] = result
        except Exception as exc:
            st.error(friendly_error(exc))
            return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
